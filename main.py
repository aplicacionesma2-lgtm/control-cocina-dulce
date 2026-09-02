import datetime
import io
import os
import openpyxl
from openpyxl.drawing.image import Image as OpenPyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
from PIL import Image as PILImage
import streamlit as st

st.set_page_config(
    page_title="Control de Proceso - Cocina Dulce",
    page_icon="📋",
    layout="wide",
)

# Inicializar historial en la memoria local de la sesión
if "historial" not in st.session_state:
    st.session_state.historial = []

FILE_PATH = "CONTROL PROCESOS COCINA DULCE.xlsx"

PRODUCTOS_DEFAULT = [
    "GLACE DE KEKE LIMÓN",
    "COMPOTA DE MUFFINS DE MANZANA",
    "COMPOTA DE KEKE GINGER",
    "COMPOTA DE CROCANTE DE MANZANA",
    "RELLENO PIE DE LIMON STBX",
    "ALMIBAR TORTA DE CHOCOLATE",
    "CREMA DE QUESO",
    "CREMA DE QUESO DE KEKES",
    "FUDGE",
    "FUDGE SILVESTRE",
    "FUDGE LIQUIDO PARA TIENDA",
    "GANACHE AGENDA",
    "MERMELADA MIXTA",
    "ZANAHORIA RALLADA",
    "ALMIBAR RED VELVET",
    "TRUFA DE CHOCOLATE 2023",
    "MANÁ",
    "MANJAR LIQUIDO",
    "MANJAR SILVESTRE",
    "MANJAR CASERO",
    "CHOCOLATE RALLADO BITTER",
    "CHOCOLATE RALLADO BLANCO MA",
    "RELLENO DE PIE DE LIMON MARIA",
    "ALMIBAR TRES LECHES VAINILLA",
    "ALMIBAR TRES LECHES CHOCOLATE",
]

EQUIPOS_DEFAULT = [
    "BATIDORA",
    "MARMITA",
    "HORNO RATIONAL",
    "ROBOTCOUPE",
    "NO APLICA",
    "OTRO",
]

RESPONSABLES_DEFAULT = [
    "CINDY VILLA",
    "JHOJAN VENTURA",
    "DANIEL CUADROS",
    "CESAR VILCHEZ",
    "LUCIA CHIQUIRENI",
    "JENNY RIVAS",
    "OTRO",
]


@st.cache_data(ttl=60)
def cargar_catalogos():
    productos = PRODUCTOS_DEFAULT
    equipos = EQUIPOS_DEFAULT
    responsables = RESPONSABLES_DEFAULT

    if os.path.exists(FILE_PATH):
        try:
            xls = pd.ExcelFile(FILE_PATH)
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(FILE_PATH, sheet_name=sheet_name)
                for idx, row in df.iterrows():
                    row_vals = [str(v).strip() for v in row.values if pd.notna(v)]
                    if "PRODUCTO" in row_vals and "EQUIPO UTILIZADO" in row_vals:
                        header_idx = idx
                        df_data = pd.read_excel(
                            FILE_PATH,
                            sheet_name=sheet_name,
                            skiprows=header_idx + 1,
                        )
                        df_data.columns = [
                            str(c).strip() for c in df.iloc[header_idx].values
                        ]

                        if "PRODUCTO" in df_data.columns:
                            prods = (
                                df_data["PRODUCTO"]
                                .dropna()
                                .astype(str)
                                .str.strip()
                                .unique()
                                .tolist()
                            )
                            if prods:
                                productos = prods

                        if "EQUIPO UTILIZADO" in df_data.columns:
                            eqs = (
                                df_data["EQUIPO UTILIZADO"]
                                .dropna()
                                .astype(str)
                                .str.strip()
                                .unique()
                                .tolist()
                            )
                            if eqs:
                                equipos = list(dict.fromkeys(eqs + ["OTRO"]))

                        if "RESPONSABLE" in df_data.columns:
                            resps = (
                                df_data["RESPONSABLE"]
                                .dropna()
                                .astype(str)
                                .str.strip()
                                .unique()
                                .tolist()
                            )
                            if resps:
                                responsables = list(
                                    dict.fromkeys(resps + ["OTRO"])
                                )
                        break
        except Exception:
            pass

    return productos, equipos, responsables


lista_productos, lista_equipos, lista_responsables = cargar_catalogos()

st.title("📋 FORMATO CONTROL DE PROCESO COCINA DULCE")
st.markdown("---")

with st.form("form_control_proceso", clear_on_submit=False):
    st.subheader("1. Información General del Proceso")
    col1, col2, col3, col4 = st.columns(4)

    with col1:
        fecha_p = st.date_input(
            "F.P (Fecha de Producción)", value=datetime.date.today()
        )
    with col2:
        lote = st.text_input(
            "LOTE", placeholder="Ej. L-20260901", key="input_lote"
        )
    with col3:
        batch = st.text_input("BATCH", placeholder="Ej. B-01", key="input_batch")
    with col4:
        responsable_sel = st.selectbox(
            "RESPONSABLE", options=lista_responsables
        )
        if responsable_sel == "OTRO":
            responsable_text = st.text_input(
                "Especifique Responsable", placeholder="Nombre completo"
            )
            responsable_final = responsable_text
        else:
            responsable_final = responsable_sel

    st.subheader("2. Selección de Producto, Equipo y Parámetros Iniciales")
    col5, col6, col_brix = st.columns([2, 2, 1])

    with col5:
        producto = st.selectbox("PRODUCTO", options=lista_productos)
    with col6:
        equipo_sel = st.selectbox("EQUIPO UTILIZADO", options=lista_equipos)
        if equipo_sel == "OTRO":
            equipo_text = st.text_input(
                "Especifique Equipo", placeholder="Nombre del equipo"
            )
            equipo_final = equipo_text
        else:
            equipo_final = equipo_sel
    with col_brix:
        brix = st.text_input(
            "BRIX (°Bx)", placeholder="Ej. 65.5", key="input_brix"
        )

    st.subheader("3. Parámetros de Control y Tiempos")
    col7, col8, col9 = st.columns(3)

    with col7:
        hora_inicio = st.time_input("HORA INICIO", value=datetime.time(8, 0))
    with col8:
        hora_termino = st.time_input("HORA TÉRMINO", value=datetime.time(8, 15))

    dt_inicio = datetime.datetime.combine(datetime.date.today(), hora_inicio)
    dt_termino = datetime.datetime.combine(datetime.date.today(), hora_termino)

    if dt_termino < dt_inicio:
        dt_termino += datetime.timedelta(days=1)

    minutos_totales = int((dt_termino - dt_inicio).total_seconds() / 60)
    tiempo_calculado = f"{minutos_totales} min"

    with col9:
        tiempo_proceso = st.text_input(
            "TIEMPO (COCCIÓN/BATIDO/ HORNEADO)",
            value=tiempo_calculado,
            disabled=True,
        )

    col10, col11 = st.columns(2)

    with col10:
        temp_equipo = st.number_input(
            "TEMPERATURA  EQUIPO (°C)", value=0.0, step=0.5, format="%.1f"
        )
    with col11:
        vel_agitador = st.text_input(
            "VELOCIDAD DEL AGITADOR (hz)/ BATIDORA/OTROS",
            placeholder="Ej. 50 Hz / V2",
        )

    st.subheader("4. Conformidad y Observaciones")
    col12, col13 = st.columns([1, 2])

    with col12:
        estado_obs = st.radio(
            "ESTADO",
            options=["CONFORME", "NO CONFORME", "OTRO (Texto Libre)"],
        )

    with col13:
        if estado_obs == "OTRO (Texto Libre)":
            obs_detalle = st.text_area(
                "OBSERVACIÓN",
                placeholder="Escriba aquí la observación personalizada...",
            )
            observacion_final = obs_detalle
        else:
            obs_adicional = st.text_input(
                "Detalle / Comentario adicional (Opcional)",
                placeholder="Escriba detalles si aplica...",
            )
            observacion_final = (
                f"{estado_obs} - {obs_adicional}".strip(" -")
                if obs_adicional
                else estado_obs
            )

    st.markdown("---")
    btn_guardar = st.form_submit_button(
        "💾 Guardar Registro Local", use_container_width=True
    )

if btn_guardar:
    nuevo_registro = {
        "PRODUCTO": producto,
        "F.P": fecha_p.strftime("%Y-%m-%d"),
        "LOTE": lote,
        "BATCH": batch,
        "EQUIPO UTILIZADO": equipo_final,
        "BRIX (°Bx)": brix,
        "HORA INICIO": hora_inicio.strftime("%H:%M"),
        "TEMPERATURA  EQUIPO (°C)": float(temp_equipo),
        "TIEMPO (COCCIÓN/BATIDO/ HORNEADO)": tiempo_calculado,
        "VELOCIDAD DEL AGITADOR (hz)/ BATIDORA/OTROS": vel_agitador,
        "HORA TÉRMINO": hora_termino.strftime("%H:%M"),
        "RESPONSABLE": responsable_final,
        "OBSERVACIÓN": observacion_final,
    }
    st.session_state.historial.append(nuevo_registro)
    st.success("✅ Registro guardado en la sesión local.")


# Generar Excel con Formato MA-FR-109 y Logo
def generar_excel_formateado(df_datos):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control Procesos"

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    fill_header = PatternFill(
        start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
    )
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_title = Font(name="Calibri", size=13, bold=True)

    # Combinación de Celdas de la Cabecera
    ws.merge_cells("A1:B3")
    ws.merge_cells("C1:K3")
    ws.merge_cells("L1:M3")

    # Título y Código MA-FR-109
    ws["C1"] = "FORMATO CONTROL DE PROCESO COCINA DULCE"
    ws["C1"].font = font_title
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")

    ws["L1"] = "MA-FR- 109"
    ws["L1"].font = font_bold
    ws["L1"].alignment = Alignment(horizontal="center", vertical="center")

    # Buscar logo en la carpeta local (incluye formato .jfif)
    logo_path = None
    for posible_logo in [
        "logo.jfif",
        "logo.png",
        "LOGOMA.jfif",
        "logo.jpg",
        "logo.jpeg",
        "LOGOMA.png",
    ]:
        if os.path.exists(posible_logo):
            logo_path = posible_logo
            break

    if logo_path:
        try:
            img = PILImage.open(logo_path).convert("RGB")
            img.thumbnail((110, 50))
            img_temp_path = "temp_logo_local.png"
            img.save(img_temp_path, "PNG")

            excel_img = OpenPyxlImage(img_temp_path)
            excel_img.anchor = "A1"
            ws.add_image(excel_img)
        except Exception:
            ws["A1"] = "MA"
            ws["A1"].alignment = Alignment(
                horizontal="center", vertical="center"
            )
    else:
        ws["A1"] = "MA"
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # Aplicar Bordes a la Cabecera
    for row in ws["A1:M3"]:
        for cell in row:
            cell.border = thin_border

    # Encabezados de Tabla (Fila 5)
    headers = list(df_datos.columns)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=header)
        cell.font = font_bold
        cell.fill = fill_header
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = thin_border

    # Escribir Datos
    for row_idx, row_data in enumerate(df_datos.values, start=6):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

    # Autoajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


st.markdown("---")
st.subheader("📊 Historial de Registros Local")

if st.session_state.historial:
    df_historial = pd.DataFrame(st.session_state.historial)
    st.dataframe(df_historial, use_container_width=True)

    excel_data = generar_excel_formateado(df_historial)

    st.download_button(
        label="📥 Descargar Excel Formateado (MA-FR-109 + Logo)",
        data=excel_data,
        file_name=f"Control_Cocina_Dulce_MA-FR-109_{datetime.date.today()}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
else:
    st.info(
        "Ingresa datos en el formulario y haz clic en 'Guardar Registro Local' para poder probar la descarga del Excel formateado."
    )